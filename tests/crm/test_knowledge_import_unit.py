"""Юнит-тесты нарезки текста и FileReader для импорта знаний (без HTTP и воркера)."""

from __future__ import annotations

from io import BytesIO

import fitz
import openpyxl
import pytest
from docx import Document

from core.files.checksum import compute_content_checksum_sha256
from core.files.reader import FileReader
from core.files.reader.models import FileReadKind
from core.utils.knowledge_text_split import (
    MAX_IMPORT_TEXT_CHARS,
    split_knowledge_text,
    validate_chunk_max_chars,
)


def test_validate_chunk_max_chars_bounds() -> None:
    validate_chunk_max_chars(2000)
    validate_chunk_max_chars(500_000)
    with pytest.raises(ValueError, match="chunk_max_chars"):
        validate_chunk_max_chars(1999)
    with pytest.raises(ValueError, match="chunk_max_chars"):
        validate_chunk_max_chars(500_001)


def test_split_empty_raises() -> None:
    with pytest.raises(ValueError, match="пуст"):
        split_knowledge_text("  \n\t  ")


def test_split_type_error() -> None:
    with pytest.raises(TypeError, match="str"):
        split_knowledge_text(None)  # type: ignore[arg-type]


def test_split_too_long_raises() -> None:
    huge = "a" * (MAX_IMPORT_TEXT_CHARS + 1)
    with pytest.raises(ValueError, match="превышает лимит"):
        split_knowledge_text(huge)


def test_split_single_chunk_trim() -> None:
    chunks = split_knowledge_text("  alpha beta  ", chunk_max_chars=50_000)
    assert chunks == ["alpha beta"]


def test_split_by_headings_two_sections() -> None:
    text = "# Раздел A\n\nУникальный маркер AAA\n\n## Раздел B\n\nУникальный маркер BBB"
    chunks = split_knowledge_text(text, split_by_headings=True, chunk_max_chars=50_000)
    assert len(chunks) == 2
    assert "AAA" in chunks[0]
    assert "BBB" in chunks[1]


def test_split_oversize_part_slices() -> None:
    body = "x" * 5500
    chunks = split_knowledge_text(body, chunk_max_chars=2000)
    assert len(chunks) == 3
    assert sum(len(c) for c in chunks) == 5500


@pytest.mark.asyncio
async def test_file_reader_markdown_bytes() -> None:
    raw = "# Заголовок\n\nТело KN_UNIT_MD_442".encode("utf-8")
    reader = FileReader()
    res = await reader.read(raw, file_name="doc.md")
    assert res.detected_kind == FileReadKind.TEXT
    joined = "\n".join(p.text for p in res.pages)
    assert "KN_UNIT_MD_442" in joined
    assert res.source_checksum == compute_content_checksum_sha256(raw)


@pytest.mark.asyncio
async def test_file_reader_xlsx_bytes() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws["A1"] = "KN_UNIT_XLSX_MARKER"
    buf = BytesIO()
    wb.save(buf)
    raw = buf.getvalue()

    reader = FileReader()
    res = await reader.read(raw, file_name="t.xlsx")
    assert res.detected_kind == FileReadKind.SPREADSHEET
    joined = "\n".join(p.text for p in res.pages)
    assert "KN_UNIT_XLSX_MARKER" in joined


@pytest.mark.asyncio
async def test_file_reader_docx_bytes() -> None:
    doc = Document()
    doc.add_paragraph("KN_UNIT_DOCX_MARKER")
    buf = BytesIO()
    doc.save(buf)
    raw = buf.getvalue()

    reader = FileReader()
    res = await reader.read(raw, file_name="t.docx")
    assert res.detected_kind == FileReadKind.OFFICE
    joined = "\n".join(p.text for p in res.pages)
    assert "KN_UNIT_DOCX_MARKER" in joined


@pytest.mark.asyncio
async def test_file_reader_pdf_bytes(tmp_path) -> None:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "KN_UNIT_PDF_MARKER")
    out = tmp_path / "u.pdf"
    doc.save(out)
    doc.close()
    raw = out.read_bytes()

    reader = FileReader()
    res = await reader.read(raw, file_name="u.pdf")
    assert res.detected_kind == FileReadKind.PDF
    joined = "\n".join(p.text for p in res.pages)
    assert "KN_UNIT_PDF_MARKER" in joined
