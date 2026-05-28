"""
E2E импорта базы знаний: реальный TaskIQ worker, S3/MinIO, парсинг файлов, откат.

LLM мокируется только в сценарии mode=graph (mock_llm_redis), как в test_05_ai_analysis.
"""

from __future__ import annotations

import json
import uuid
from collections.abc import Awaitable, Callable
from io import BytesIO
from pathlib import Path
from typing import cast

import fitz
import openpyxl
import pytest
from docx import Document
from httpx import AsyncClient, Response

from tests.crm.e2e._json_helpers import json_object, object_dict, object_list, object_str
from tests.crm.knowledge_import_helpers import (
    combined_entity_blob,
    crm_upload_bytes,
    fetch_entity_texts,
    rollback_task,
    wait_task_terminal,
)
from tests.fixtures.crm_test_setup import ensure_entity_type_in_namespace

MockLlmRedisFactory = Callable[[list[object]], Awaitable[None]]

_META: dict[str, object] = {"dates_mentioned": [], "places_mentioned": [], "key_topics": []}

pytestmark = [
    pytest.mark.real_taskiq,
    pytest.mark.timeout(120, func_only=True),
]


def _http_json(response: Response) -> dict[str, object]:
    return json_object(cast(object, response.json()))


def _string_list(value: object) -> list[str]:
    if not isinstance(value, list):
        return []
    strings: list[str] = []
    typed_list = cast(list[object], value)
    for raw_item in typed_list:
        if isinstance(raw_item, str):
            strings.append(raw_item)
    return strings


def _task_data(row: dict[str, object]) -> dict[str, object]:
    data_raw = row.get("data")
    if isinstance(data_raw, dict):
        return object_dict(cast(object, data_raw), field="data")
    return {}


def _created_entity_ids(row: dict[str, object]) -> list[str]:
    return _string_list(_task_data(row).get("created_entity_ids"))


async def _wait_task_row(
    crm_client: AsyncClient,
    headers: dict[str, str],
    task_id: str,
) -> dict[str, object]:
    row = await wait_task_terminal(crm_client, headers, task_id)
    return json_object(cast(object, row))


def _test_namespace(unique_id: str) -> str:
    return f"g_{unique_id}_{uuid.uuid4().hex[:8]}"


def _xlsx_bytes(marker: str) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws["A1"] = marker
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx_bytes(marker: str) -> bytes:
    doc = Document()
    _ = doc.add_paragraph(marker)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pdf_bytes(tmp_path: Path, marker: str) -> bytes:
    doc = fitz.open()  # pyright: ignore[reportCallIssue, reportUnknownVariableType]
    page = doc.new_page()  # pyright: ignore[reportUnknownMemberType, reportUnknownVariableType]
    page.insert_text((72, 72), marker)  # pyright: ignore[reportUnknownMemberType]
    out = tmp_path / "kn_e2e.pdf"
    doc.save(out)  # pyright: ignore[reportUnknownMemberType]
    doc.close()  # pyright: ignore[reportUnknownMemberType]
    return out.read_bytes()


async def _start_notes_import(
    crm_client: AsyncClient,
    headers: dict[str, str],
    *,
    namespace: str,
    source_text: str | None = None,
    source_file_ids: list[str] | None = None,
    split_by_headings: bool = False,
    chunk_max_chars: int = 50_000,
) -> str:
    await _ensure_import_namespace(crm_client, headers, namespace)

    body: dict[str, object] = {
        "namespace": namespace,
        "mode": "notes_only",
        "split_by_headings": split_by_headings,
        "chunk_max_chars": chunk_max_chars,
    }
    if source_text is not None:
        body["source_text"] = source_text
    if source_file_ids is not None:
        body["source_file_ids"] = source_file_ids
    response = await crm_client.post(
        "/crm/api/v1/tasks/knowledge-import",
        json=body,
        headers=headers,
    )
    if response.status_code != 202:
        raise AssertionError(f"start import: {response.status_code} {response.text}")
    payload = _http_json(response)
    return object_str(payload.get("task_id"), field="task_id").strip()


async def _ensure_import_namespace(
    crm_client: AsyncClient,
    headers: dict[str, str],
    namespace: str,
) -> None:
    create_ns = await crm_client.post(
        "/crm/api/v1/namespaces",
        json={"name": namespace, "template_id": "sales"},
        headers=headers,
    )
    if create_ns.status_code not in (201, 409):
        raise AssertionError(f"create namespace: {create_ns.status_code} {create_ns.text}")
    editability = await crm_client.get(
        f"/crm/api/v1/namespaces/{namespace}/editability",
        headers=headers,
    )
    if editability.status_code != 200:
        raise AssertionError(f"namespace editability: {editability.status_code} {editability.text}")
    edit_payload = _http_json(editability)
    allowed = _string_list(edit_payload.get("current_allowed_type_ids"))
    target_allowed = sorted({*allowed, "note", "meeting", "task"})
    update_ns = await crm_client.put(
        f"/crm/api/v1/namespaces/{namespace}",
        json={"allowed_type_ids": target_allowed},
        headers=headers,
    )
    if update_ns.status_code != 200:
        raise AssertionError(f"namespace update: {update_ns.status_code} {update_ns.text}")
    for system_type in ("note", "meeting", "task"):
        _ = await ensure_entity_type_in_namespace(
            crm_client,
            headers,
            system_type,
            system_type,
            namespace,
        )


async def _assert_blob_then_rollback(
    crm_client: AsyncClient,
    headers: dict[str, str],
    *,
    task_id: str,
    must_contain: str,
) -> None:
    row = await _wait_task_row(crm_client, headers, task_id)
    assert row.get("status") == "completed"
    ids = _created_entity_ids(row)
    assert len(ids) >= 1
    blob = combined_entity_blob(await fetch_entity_texts(crm_client, headers, ids))
    assert must_contain in blob
    _ = await rollback_task(crm_client, headers, task_id)
    final = await crm_client.get(
        f"/crm/api/v1/tasks/{task_id}",
        headers=headers,
    )
    assert final.status_code == 200
    assert _http_json(final).get("status") == "rolled_back"
    for eid in ids:
        er = await crm_client.get(f"/crm/api/v1/entities/{eid}", headers=headers)
        assert er.status_code == 404, er.text


class TestKnowledgeImportE2E:
    @pytest.mark.asyncio
    async def test_notes_only_inline_text(
        self,
        crm_client: AsyncClient,
        crm_worker: object,
        unique_id: str,
        auth_headers_system: dict[str, str],
    ) -> None:
        _ = crm_worker
        ns = _test_namespace(unique_id)
        marker = f"KN_E2E_INLINE_{unique_id}"
        task_id = await _start_notes_import(
            crm_client,
            auth_headers_system,
            namespace=ns,
            source_text=f"Заголовок\n\n{marker}\n",
        )
        await _assert_blob_then_rollback(
            crm_client,
            auth_headers_system,
            task_id=task_id,
            must_contain=marker,
        )

    @pytest.mark.asyncio
    async def test_notes_only_markdown_file(
        self,
        crm_client: AsyncClient,
        crm_worker: object,
        unique_id: str,
        auth_headers_system: dict[str, str],
    ) -> None:
        _ = crm_worker
        ns = _test_namespace(unique_id)
        marker = f"KN_E2E_MD_{unique_id}"
        raw = f"# Раздел\n\n{marker}\n".encode("utf-8")
        fid = await crm_upload_bytes(crm_client, auth_headers_system, f"{unique_id}.md", raw)
        task_id = await _start_notes_import(
            crm_client,
            auth_headers_system,
            namespace=ns,
            source_file_ids=[fid],
        )
        await _assert_blob_then_rollback(
            crm_client,
            auth_headers_system,
            task_id=task_id,
            must_contain=marker,
        )

    @pytest.mark.asyncio
    async def test_notes_only_txt_file(
        self,
        crm_client: AsyncClient,
        crm_worker: object,
        unique_id: str,
        auth_headers_system: dict[str, str],
    ) -> None:
        _ = crm_worker
        ns = _test_namespace(unique_id)
        marker = f"KN_E2E_TXT_{unique_id}"
        raw = f"plain\n{marker}\n".encode("utf-8")
        fid = await crm_upload_bytes(crm_client, auth_headers_system, f"{unique_id}.txt", raw)
        task_id = await _start_notes_import(
            crm_client,
            auth_headers_system,
            namespace=ns,
            source_file_ids=[fid],
        )
        await _assert_blob_then_rollback(
            crm_client,
            auth_headers_system,
            task_id=task_id,
            must_contain=marker,
        )

    @pytest.mark.asyncio
    async def test_notes_only_csv_file(
        self,
        crm_client: AsyncClient,
        crm_worker: object,
        unique_id: str,
        auth_headers_system: dict[str, str],
    ) -> None:
        _ = crm_worker
        ns = _test_namespace(unique_id)
        marker = f"KN_E2E_CSV_{unique_id}"
        raw = f"col1,col2\n1,{marker}\n".encode("utf-8")
        fid = await crm_upload_bytes(crm_client, auth_headers_system, f"{unique_id}.csv", raw)
        task_id = await _start_notes_import(
            crm_client,
            auth_headers_system,
            namespace=ns,
            source_file_ids=[fid],
        )
        await _assert_blob_then_rollback(
            crm_client,
            auth_headers_system,
            task_id=task_id,
            must_contain=marker,
        )

    @pytest.mark.asyncio
    async def test_notes_only_xlsx_file(
        self,
        crm_client: AsyncClient,
        crm_worker: object,
        unique_id: str,
        auth_headers_system: dict[str, str],
    ) -> None:
        _ = crm_worker
        ns = _test_namespace(unique_id)
        marker = f"KN_E2E_XLSX_{unique_id}"
        raw = _xlsx_bytes(marker)
        fid = await crm_upload_bytes(crm_client, auth_headers_system, f"{unique_id}.xlsx", raw)
        task_id = await _start_notes_import(
            crm_client,
            auth_headers_system,
            namespace=ns,
            source_file_ids=[fid],
        )
        await _assert_blob_then_rollback(
            crm_client,
            auth_headers_system,
            task_id=task_id,
            must_contain=marker,
        )

    @pytest.mark.asyncio
    async def test_notes_only_docx_file(
        self,
        crm_client: AsyncClient,
        crm_worker: object,
        unique_id: str,
        auth_headers_system: dict[str, str],
    ) -> None:
        _ = crm_worker
        ns = _test_namespace(unique_id)
        marker = f"KN_E2E_DOCX_{unique_id}"
        raw = _docx_bytes(marker)
        fid = await crm_upload_bytes(crm_client, auth_headers_system, f"{unique_id}.docx", raw)
        task_id = await _start_notes_import(
            crm_client,
            auth_headers_system,
            namespace=ns,
            source_file_ids=[fid],
        )
        await _assert_blob_then_rollback(
            crm_client,
            auth_headers_system,
            task_id=task_id,
            must_contain=marker,
        )

    @pytest.mark.asyncio
    async def test_notes_only_pdf_file(
        self,
        crm_client: AsyncClient,
        crm_worker: object,
        unique_id: str,
        auth_headers_system: dict[str, str],
        tmp_path: Path,
    ) -> None:
        _ = crm_worker
        ns = _test_namespace(unique_id)
        marker = f"KN_E2E_PDF_{unique_id}"
        raw = _pdf_bytes(tmp_path, marker)
        fid = await crm_upload_bytes(crm_client, auth_headers_system, f"{unique_id}.pdf", raw)
        task_id = await _start_notes_import(
            crm_client,
            auth_headers_system,
            namespace=ns,
            source_file_ids=[fid],
        )
        await _assert_blob_then_rollback(
            crm_client,
            auth_headers_system,
            task_id=task_id,
            must_contain=marker,
        )

    @pytest.mark.asyncio
    async def test_notes_only_two_files_and_inline(
        self,
        crm_client: AsyncClient,
        crm_worker: object,
        unique_id: str,
        auth_headers_system: dict[str, str],
    ) -> None:
        _ = crm_worker
        ns = _test_namespace(unique_id)
        m1 = f"KN_E2E_M1_{unique_id}"
        m2 = f"KN_E2E_M2_{unique_id}"
        m0 = f"KN_E2E_M0_{unique_id}"
        f1 = await crm_upload_bytes(
            crm_client,
            auth_headers_system,
            f"a_{unique_id}.txt",
            f"file1\n{m1}\n".encode("utf-8"),
        )
        f2 = await crm_upload_bytes(
            crm_client,
            auth_headers_system,
            f"b_{unique_id}.txt",
            f"file2\n{m2}\n".encode("utf-8"),
        )
        task_id = await _start_notes_import(
            crm_client,
            auth_headers_system,
            namespace=ns,
            source_text=f"inline\n{m0}\n",
            source_file_ids=[f1, f2],
        )
        row = await _wait_task_row(crm_client, auth_headers_system, task_id)
        assert row.get("status") == "completed"
        ids = _created_entity_ids(row)
        blob = combined_entity_blob(await fetch_entity_texts(crm_client, auth_headers_system, ids))
        assert m0 in blob and m1 in blob and m2 in blob
        _ = await rollback_task(crm_client, auth_headers_system, task_id)

    @pytest.mark.asyncio
    async def test_notes_only_split_by_headings_two_notes(
        self,
        crm_client: AsyncClient,
        crm_worker: object,
        unique_id: str,
        auth_headers_system: dict[str, str],
    ) -> None:
        _ = crm_worker
        ns = _test_namespace(unique_id)
        a = f"KN_E2E_HA_{unique_id}"
        b = f"KN_E2E_HB_{unique_id}"
        text = f"# Part one\n\n{a}\n\n## Part two\n\n{b}\n"
        task_id = await _start_notes_import(
            crm_client,
            auth_headers_system,
            namespace=ns,
            source_text=text,
            split_by_headings=True,
        )
        row = await _wait_task_row(crm_client, auth_headers_system, task_id)
        assert _task_data(row).get("notes_created_count") == 2
        ids = _created_entity_ids(row)
        blob = combined_entity_blob(await fetch_entity_texts(crm_client, auth_headers_system, ids))
        assert a in blob and b in blob
        _ = await rollback_task(crm_client, auth_headers_system, task_id)

    @pytest.mark.asyncio
    async def test_review_list_and_complete_then_rollback(
        self,
        crm_client: AsyncClient,
        crm_worker: object,
        unique_id: str,
        auth_headers_system: dict[str, str],
    ) -> None:
        _ = crm_worker
        ns = _test_namespace(unique_id)
        marker = f"KN_E2E_REV_{unique_id}"
        task_id = await _start_notes_import(
            crm_client,
            auth_headers_system,
            namespace=ns,
            source_text=f"text\n{marker}\n",
        )
        row = await _wait_task_row(crm_client, auth_headers_system, task_id)
        assert row.get("status") == "completed"
        list_r = await crm_client.get(
            f"/crm/api/v1/tasks/{task_id}/created-entities",
            headers=auth_headers_system,
        )
        assert list_r.status_code == 200
        listed = _http_json(list_r)
        assert listed.get("task_id") == task_id
        assert len(object_list(listed.get("entities"))) >= 1
        done = await crm_client.post(
            f"/crm/api/v1/tasks/{task_id}/review-complete",
            headers=auth_headers_system,
        )
        assert done.status_code == 200
        done_data = _task_data(_http_json(done))
        assert done_data.get("review_completed_at") is not None
        _ = await rollback_task(crm_client, auth_headers_system, task_id)

    @pytest.mark.asyncio
    async def test_graph_mode_single_chunk_mock_llm(
        self,
        crm_client: AsyncClient,
        crm_worker: object,
        unique_id: str,
        auth_headers_system: dict[str, str],
        mock_llm_redis: MockLlmRedisFactory,
    ) -> None:
        _ = crm_worker
        ns = _test_namespace(unique_id)
        chunk_marker = f"KN_E2E_GRAPH_CHUNK_{unique_id}"
        note_title = f"Заметка граф {unique_id}"
        task_name = f"Задача граф {unique_id}"
        payload: dict[str, object] = {
            "note": {
                "entity_type": "note",
                "name": note_title,
                "description": f"Описание {chunk_marker}",
            },
            "entities": [
                {
                    "entity_type": "task",
                    "name": task_name,
                    "description": "Задача из импорта графом",
                    "attributes": {"origin": "knowledge_import"},
                },
            ],
            "relationships": [
                {
                    "source_type": "note",
                    "source_name": note_title,
                    "target_type": "task",
                    "target_name": task_name,
                    "relationship_type": "mentions",
                    "weight": 1.0,
                    "confidence": 0.9,
                },
            ],
            "metadata": _META,
        }
        await mock_llm_redis(
            [
                {"type": "text", "content": json.dumps(payload)},
                {"type": "text", "content": json.dumps(payload)},
            ]
        )
        await _ensure_import_namespace(crm_client, auth_headers_system, ns)
        body: dict[str, object] = {
            "namespace": ns,
            "mode": "graph",
            "source_text": f"Текст для анализа. {chunk_marker}",
            "extract_entity_types": ["task"],
            "split_by_headings": False,
            "chunk_max_chars": 50_000,
        }
        response = await crm_client.post(
            "/crm/api/v1/tasks/knowledge-import",
            json=body,
            headers=auth_headers_system,
        )
        assert response.status_code == 202, response.text
        task_id = object_str(_http_json(response).get("task_id"), field="task_id")
        row = await _wait_task_row(crm_client, auth_headers_system, task_id)
        assert row.get("status") == "completed"
        task_data = _task_data(row)
        notes_created = task_data.get("notes_created_count")
        entities_created = task_data.get("entities_created_count")
        assert isinstance(notes_created, int) and notes_created >= 1
        assert isinstance(entities_created, int) and entities_created >= 1
        ids = _created_entity_ids(row)
        blob = combined_entity_blob(await fetch_entity_texts(crm_client, auth_headers_system, ids))
        assert chunk_marker in blob
        assert task_name in blob
        _ = await rollback_task(crm_client, auth_headers_system, task_id)

    @pytest.mark.asyncio
    async def test_graph_mode_meeting_stored_as_note_for_notes_ui(
        self,
        crm_client: AsyncClient,
        crm_worker: object,
        unique_id: str,
        auth_headers_system: dict[str, str],
        mock_llm_redis: MockLlmRedisFactory,
    ) -> None:
        _ = crm_worker
        ns = _test_namespace(unique_id)
        chunk_marker = f"KN_E2E_MEET_NOTE_{unique_id}"
        note_title = f"Заметка контейнер {unique_id}"
        meeting_name = f"Встреча E2E {unique_id}"
        await mock_llm_redis(
            [
                {
                    "type": "text",
                    "content": json.dumps(
                        {
                            "note": {
                                "entity_type": "note",
                                "name": note_title,
                                "description": f"Текст чанка {chunk_marker}",
                            },
                            "entities": [
                                {
                                    "entity_type": "meeting",
                                    "name": meeting_name,
                                    "description": "Онлайн звонок",
                                    "attributes": {"origin": "knowledge_import"},
                                },
                            ],
                            "relationships": [
                                {
                                    "source_type": "note",
                                    "source_name": note_title,
                                    "target_type": "meeting",
                                    "target_name": meeting_name,
                                    "relationship_type": "mentions",
                                    "weight": 1.0,
                                    "confidence": 0.9,
                                },
                            ],
                            "metadata": _META,
                        }
                    ),
                }
            ]
        )
        await _ensure_import_namespace(crm_client, auth_headers_system, ns)
        body: dict[str, object] = {
            "namespace": ns,
            "mode": "graph",
            "source_text": f"Контент импорта. {chunk_marker}",
            "extract_entity_types": ["meeting"],
            "split_by_headings": False,
            "chunk_max_chars": 50_000,
        }
        response = await crm_client.post(
            "/crm/api/v1/tasks/knowledge-import",
            json=body,
            headers=auth_headers_system,
        )
        assert response.status_code == 202, response.text
        task_id = object_str(_http_json(response).get("task_id"), field="task_id")
        row = await _wait_task_row(crm_client, auth_headers_system, task_id)
        assert row.get("status") == "completed"
        list_r = await crm_client.post(
            "/crm/api/v1/entities/query",
            json={
                "namespace": ns,
                "entity_type": "note",
                "limit": 50,
                "search_mode": "hybrid",
            },
            headers=auth_headers_system,
        )
        assert list_r.status_code == 200, list_r.text
        query_payload = _http_json(list_r)
        items_raw = query_payload.get("items")
        items = object_list(items_raw if items_raw is not None else query_payload)
        assert any(
            object_dict(entity, field="entity").get("entity_type") == "note"
            and object_dict(entity, field="entity").get("entity_subtype") in ("meeting", None, "")
            for entity in items
        ), items
        _ = await rollback_task(crm_client, auth_headers_system, task_id)
