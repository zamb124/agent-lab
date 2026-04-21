"""Тесты чтения устаревших форматов .xls и .doc через FileReader."""

from __future__ import annotations

import shutil
import subprocess
from io import BytesIO
from pathlib import Path

import pytest
import xlwt

from core.files.reader import FileReader, FileReadError
from core.files.reader.models import FileReadKind


def _make_xls_bytes(sheet_name: str, rows: list[list[str]]) -> bytes:
    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet(sheet_name)
    for row_idx, row in enumerate(rows):
        for col_idx, value in enumerate(row):
            ws.write(row_idx, col_idx, value)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_docx_bytes(text: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ─────────────────────────────────────────── XLS ────────────────────────────


@pytest.mark.asyncio
async def test_read_xls_single_sheet(tmp_path: Path) -> None:
    rows = [
        ["Товар", "Количество", "Цена"],
        ["Яблоки", "100", "50.5"],
        ["Груши", "200", "75.0"],
    ]
    raw = _make_xls_bytes("Склад", rows)

    reader = FileReader()
    result = await reader.read(raw, file_name="report.xls")

    assert result.detected_kind == FileReadKind.SPREADSHEET
    assert result.page_count == 1
    assert result.pages[0].label == "Склад"

    text = result.pages[0].text
    assert "Товар" in text
    assert "Яблоки" in text
    assert "Груши" in text
    assert "100" in text


@pytest.mark.asyncio
async def test_read_xls_multiple_sheets(tmp_path: Path) -> None:
    wb = xlwt.Workbook(encoding="utf-8")

    ws1 = wb.add_sheet("Январь")
    ws1.write(0, 0, "Строка-Январь-Маркер")

    ws2 = wb.add_sheet("Февраль")
    ws2.write(0, 0, "Строка-Февраль-Маркер")

    buf = BytesIO()
    wb.save(buf)
    raw = buf.getvalue()

    reader = FileReader()
    result = await reader.read(raw, file_name="months.xls")

    assert result.detected_kind == FileReadKind.SPREADSHEET
    assert result.page_count == 2

    labels = [p.label for p in result.pages]
    assert "Январь" in labels
    assert "Февраль" in labels

    all_text = "\n".join(p.text for p in result.pages)
    assert "Строка-Январь-Маркер" in all_text
    assert "Строка-Февраль-Маркер" in all_text


@pytest.mark.asyncio
async def test_read_xls_numeric_values(tmp_path: Path) -> None:
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Data")
    ws.write(0, 0, 42)
    ws.write(0, 1, 3.14)
    ws.write(1, 0, "текст")
    buf = BytesIO()
    wb.save(buf)

    reader = FileReader()
    result = await reader.read(buf.getvalue(), file_name="nums.xls")

    assert result.detected_kind == FileReadKind.SPREADSHEET
    text = result.pages[0].text
    assert "42" in text
    assert "текст" in text


@pytest.mark.asyncio
async def test_read_xls_empty_raises() -> None:
    """Пустой XLS без данных поднимает FileReadError."""
    wb = xlwt.Workbook()
    wb.add_sheet("Empty")
    buf = BytesIO()
    wb.save(buf)

    reader = FileReader()
    with pytest.raises(FileReadError):
        await reader.read(buf.getvalue(), file_name="empty.xls")


# ─────────────────────────────────────────── DOC ────────────────────────────


_FIXTURES_DIR = Path(__file__).parent / "fixtures"
_SAMPLE_DOC = _FIXTURES_DIR / "sample.doc"


@pytest.mark.asyncio
@pytest.mark.timeout(30)
async def test_read_doc_via_antiword() -> None:
    assert shutil.which("antiword"), "Для чтения .doc нужен исполняемый antiword в PATH"
    assert _SAMPLE_DOC.is_file(), f"Фикстура .doc не найдена: {_SAMPLE_DOC}"
    reader = FileReader()
    result = await reader.read(_SAMPLE_DOC)

    assert result.detected_kind == FileReadKind.OFFICE
    assert result.page_count >= 1
    all_text = "\n".join(p.text for p in result.pages)
    assert len(all_text.strip()) > 0


@pytest.mark.asyncio
async def test_read_doc_without_antiword_raises_clear_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr("shutil.which", lambda _: None)

    doc_path = tmp_path / "source.doc"
    doc_path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512)

    reader = FileReader()
    with pytest.raises(FileReadError, match="antiword"):
        await reader.read(doc_path)
